#!/usr/bin/env python3
"""Generate the VitalAI x Luli Fitness PowerPoint deck.

Mirrors the reveal.js deck (docs/luli-fitness-presentation.html) as an
editable 16:9 .pptx with the same dark theme and blue accent.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Palette (matches the HTML deck) ----
BG       = RGBColor(0x0A, 0x0A, 0x0A)
CARD     = RGBColor(0x16, 0x16, 0x16)
ACCENT   = RGBColor(0x4F, 0xC3, 0xF7)
WARN     = RGBColor(0xFF, 0x70, 0x43)
GREEN    = RGBColor(0x66, 0xBB, 0x6A)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
TEXT     = RGBColor(0xCC, 0xCC, 0xCC)
MUTED    = RGBColor(0x88, 0x88, 0x88)
DIM      = RGBColor(0x66, 0x66, 0x66)
WARN_BG  = RGBColor(0x1F, 0x14, 0x10)
ACC_BG   = RGBColor(0x0E, 0x1C, 0x24)
CARD_LN  = RGBColor(0x22, 0x22, 0x22)

FONT = "Arial"

EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = BG
    r.line.fill.background()
    r.shadow.inherit = False
    # send bg to back
    sp = r._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def textbox(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tf


def para(tf, runs, size, align=PP_ALIGN.LEFT, bold=False, color=TEXT,
         space_after=6, line=1.15, first=False):
    """runs: str or list of (text, color, bold) tuples."""
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    try:
        p.line_spacing = line
    except Exception:
        pass
    if isinstance(runs, str):
        runs = [(runs, color, bold)]
    for txt, c, b in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = b
        r.font.color.rgb = c; r.font.name = FONT
    return p


def card(s, l, t, w, h, fill, bar):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.fill.background(); box.shadow.inherit = False
    try:
        box.adjustments[0] = 0.04
    except Exception:
        pass
    bk = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(0.05), h)
    bk.fill.solid(); bk.fill.fore_color.rgb = bar
    bk.line.fill.background(); bk.shadow.inherit = False
    return box


def heading(s, text, color=WHITE, size=34, top=Inches(0.55)):
    tf = textbox(s, Inches(0.9), top, Inches(11.5), Inches(1.0))
    para(tf, text, size, bold=True, color=color, first=True, space_after=0)
    return tf


# =====================================================================
# SLIDE 1 — Title
# =====================================================================
s = slide()
tf = textbox(s, Inches(0), Inches(2.7), EMU_W, Inches(1.5), MSO_ANCHOR.MIDDLE)
para(tf, [("Vital", WHITE, True), ("AI", ACCENT, True)], 66, PP_ALIGN.CENTER,
     first=True, space_after=4)
tf2 = textbox(s, Inches(0), Inches(4.0), EMU_W, Inches(0.6), MSO_ANCHOR.TOP)
para(tf2, "The health coach that speaks first", 22, PP_ALIGN.CENTER, color=MUTED, first=True)
tf3 = textbox(s, Inches(0), Inches(4.6), EMU_W, Inches(0.5), MSO_ANCHOR.TOP)
para(tf3, "Luis Amancio  ·  June 2026", 14, PP_ALIGN.CENTER, color=DIM, first=True)

# =====================================================================
# SLIDE 2 — The Problem
# =====================================================================
s = slide()
heading(s, "The Problem", size=36, top=Inches(0.7))
stats = [("538M", "wearables shipped in 2024"),
         ("86%", "of users abandon diet &\nnutrition apps"),
         ("~3%", "health & fitness app\nretention at day 30")]
xs = [Inches(1.0), Inches(5.1), Inches(9.2)]
for (num, lab), x in zip(stats, xs):
    tf = textbox(s, x, Inches(2.4), Inches(3.1), Inches(1.2), MSO_ANCHOR.TOP)
    para(tf, num, 54, PP_ALIGN.CENTER, bold=True, color=ACCENT, first=True, space_after=0)
    tf2 = textbox(s, x, Inches(3.7), Inches(3.1), Inches(1.0), MSO_ANCHOR.TOP)
    para(tf2, lab, 14, PP_ALIGN.CENTER, color=MUTED, first=True, line=1.2)
tf = textbox(s, Inches(1.4), Inches(5.5), Inches(10.5), Inches(0.8), MSO_ANCHOR.TOP)
para(tf, "The data exists. What's missing is someone to turn data into conversation, and conversation into habit.",
     18, PP_ALIGN.CENTER, color=RGBColor(0xAA,0xAA,0xAA), first=True)

# =====================================================================
# SLIDE 3 — The Broken Cycle
# =====================================================================
s = slide()
heading(s, "The Broken Cycle", size=36, top=Inches(0.9))
tf = textbox(s, Inches(1.0), Inches(2.7), Inches(11.3), Inches(0.8), MSO_ANCHOR.MIDDLE)
para(tf, [("Data generated → User opens app → ", MUTED, False),
          ("“So what?”", WARN, True),
          (" → Forgets", MUTED, False)], 22, PP_ALIGN.CENTER, first=True)
tf2 = textbox(s, Inches(1.0), Inches(4.2), Inches(11.3), Inches(0.8), MSO_ANCHOR.MIDDLE)
para(tf2, [("Data generated → ", TEXT, False),
           ("AI detects context", ACCENT, True),
           (" → ", TEXT, False),
           ("Message at the right moment", ACCENT, True),
           (" → ", TEXT, False),
           ("Action", GREEN, True)], 22, PP_ALIGN.CENTER, first=True)

# =====================================================================
# SLIDE 4 — Vision quote
# =====================================================================
s = slide()
card(s, Inches(1.4), Inches(2.6), Inches(10.5), Inches(2.3), ACC_BG, ACCENT)
tf = textbox(s, Inches(2.0), Inches(2.6), Inches(9.3), Inches(2.3), MSO_ANCHOR.MIDDLE)
para(tf, "“What if your smartwatch could know you so well that it knew when to talk to you, what to say, and how to help — without you ever asking?”",
     26, PP_ALIGN.CENTER, color=RGBColor(0xE0,0xE0,0xE0), first=True, line=1.3)

# =====================================================================
# SLIDE 5 — A Day With VitalAI
# =====================================================================
s = slide()
heading(s, "A Day With VitalAI", size=32, top=Inches(0.5))
day = [("07:12", "Detected: you woke up",
        "“Good morning! You slept 6h48m, quality score 72. Your HRV is below normal — suggested breakfast: scrambled eggs with whole wheat toast and fruit.”"),
       ("12:30", "Detected: eating gesture",
        "“Eating now?” — 1 tap to confirm. Meal logged."),
       ("18:30", "Detected: workout ended",
        "“Great workout! 42min, ~380kcal. For recovery: chicken wrap with rice? Takes 15min.”"),
       ("21:45", "Detected: elevated stress",
        "“I noticed your stress is elevated. How about 5min of breathing before bed?”")]
y = Inches(1.6)
for time, title, msg in day:
    tf = textbox(s, Inches(1.0), y, Inches(1.1), Inches(1.0), MSO_ANCHOR.TOP)
    para(tf, time, 16, PP_ALIGN.LEFT, bold=True, color=ACCENT, first=True)
    box = card(s, Inches(2.2), y, Inches(10.0), Inches(1.15), CARD, CARD)
    tf2 = textbox(s, Inches(2.5), y+Inches(0.1), Inches(9.5), Inches(0.95), MSO_ANCHOR.MIDDLE)
    para(tf2, title, 13, bold=True, color=ACCENT, first=True, space_after=2)
    para(tf2, msg, 12, color=RGBColor(0xDD,0xDD,0xDD), line=1.15)
    y += Inches(1.32)

# =====================================================================
# SLIDE 6 — Three Differentiators
# =====================================================================
s = slide()
heading(s, "Three Differentiators", size=36, top=Inches(1.0))
tf = textbox(s, Inches(1.0), Inches(2.2), Inches(11.3), Inches(0.5), MSO_ANCHOR.TOP)
para(tf, "No competitor delivers all three together:", 20, PP_ALIGN.CENTER, color=TEXT, first=True)
items = ["True proactivity via biometrics",
         "Automatic meal detection",
         "Nutrition feedback loop that learns"]
y = Inches(3.3)
for i, it in enumerate(items, 1):
    tf = textbox(s, Inches(2.5), y, Inches(8.3), Inches(0.6), MSO_ANCHOR.TOP)
    para(tf, [(f"{i}.  ", ACCENT, True), (it, WHITE, False)], 24, PP_ALIGN.LEFT, first=True)
    y += Inches(0.85)

# =====================================================================
# Helper for differentiator number tag
# =====================================================================
def diff_header(s, tag, title, title_color=WHITE):
    tf = textbox(s, Inches(0.9), Inches(0.45), Inches(11.5), Inches(0.45), MSO_ANCHOR.TOP)
    para(tf, tag, 18, bold=True, color=ACCENT, first=True, space_after=0)
    tf2 = textbox(s, Inches(0.9), Inches(0.95), Inches(11.5), Inches(0.8), MSO_ANCHOR.TOP)
    para(tf2, title, 30, bold=True, color=title_color, first=True, space_after=0)

# =====================================================================
# SLIDE 7 — Differentiator 1: Proactive
# =====================================================================
s = slide()
diff_header(s, "Differentiator 1", "Proactive Coaching via Biometrics")
rows = [("Context detected", "Trigger", "AI action"),
        ("Woke up", "HR stabilized + movement", "Sleep summary + breakfast"),
        ("Workout ended", "HR returned to normal", "Recovery meal suggestion"),
        ("High stress", "HRV dropped below baseline", "Break / breathing suggestion"),
        ("Long inactivity", "90min+ without movement", "Gentle movement reminder"),
        ("Poor sleep", "Low HRV + sleep < 6h", "Adjusted suggestions for the day")]
tbl_w = Inches(11.5); tbl_h = Inches(3.4)
gtbl = s.shapes.add_table(len(rows), 3, Inches(0.9), Inches(2.0), tbl_w, tbl_h).table
gtbl.columns[0].width = Inches(3.3)
gtbl.columns[1].width = Inches(4.1)
gtbl.columns[2].width = Inches(4.1)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = gtbl.cell(ri, ci)
        cell.margin_left = Inches(0.15); cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x12,0x22,0x2A) if ri == 0 else BG
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.size = Pt(13); r.font.name = FONT
        r.font.bold = (ri == 0)
        r.font.color.rgb = ACCENT if ri == 0 else TEXT
tf = textbox(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(0.5), MSO_ANCHOR.TOP)
para(tf, "Every competitor (Oura, Fitbit, WHOOP) is reactive — the user has to open the app.",
     14, PP_ALIGN.LEFT, color=MUTED, first=True)

# =====================================================================
# Challenge-slide helper
# =====================================================================
def challenge_slide(tag, subtitle, cards, bar=WARN, fill=WARN_BG, strong=WARN):
    s = slide()
    tf = textbox(s, Inches(0.9), Inches(0.45), Inches(11.5), Inches(0.45), MSO_ANCHOR.TOP)
    para(tf, tag, 18, bold=True, color=ACCENT, first=True, space_after=0)
    tf2 = textbox(s, Inches(0.9), Inches(0.95), Inches(11.5), Inches(0.7), MSO_ANCHOR.TOP)
    para(tf2, subtitle, 24, bold=True, color=bar, first=True, space_after=0)
    n = len(cards)
    top = Inches(2.05)
    gap = Inches(0.18)
    avail = Inches(5.0)
    ch = (avail - gap*(n-1)) / n
    y = top
    for head, body in cards:
        card(s, Inches(0.9), y, Inches(11.5), ch, fill, bar)
        tf = textbox(s, Inches(1.25), y+Inches(0.12), Inches(10.9), ch-Inches(0.24), MSO_ANCHOR.MIDDLE)
        para(tf, [(head + "  ", strong, True), (body, TEXT, False)], 14, first=True, line=1.15)
        y = Emu(int(y) + int(ch) + int(gap))
    return s

# =====================================================================
# SLIDE 8 — Challenge 1
# =====================================================================
challenge_slide("Differentiator 1 — The Challenges",
    "Proactivity is a double-edged sword",
    [("Notification fatigue is the #1 risk.",
      "If notifications feel irrelevant or too frequent, users disable them and never come back. Calibrating this requires real data from real users."),
     ("Health thresholds are not universal.",
      "A resting HR of 55bpm is normal for a runner, alarming for a sedentary person. HRV of 30ms could mean stress or just be someone's baseline. The MVP uses fixed multipliers — it needs to evolve toward individual personalization."),
     ("Apple controls the platform.",
      "iOS decides when and how often background monitoring runs. If the user force-closes the app, triggers stop firing. We're building on a platform we don't control.")])

# =====================================================================
# SLIDE 9 — Differentiator 2: Meal Detection
# =====================================================================
s = slide()
diff_header(s, "Differentiator 2", "Automatic Meal Detection")
tf = textbox(s, Inches(0.9), Inches(1.75), Inches(11.5), Inches(0.4), MSO_ANCHOR.TOP)
para(tf, "The end of manual logging — the #1 reason for nutrition app abandonment",
     15, color=RGBColor(0xAA,0xAA,0xAA), first=True)
layers = [("Layer 1: Watch Gesture",
           "Accelerometer + gyroscope detect hand-to-mouth movement. Core ML model runs locally on watch.",
           "Architecture ready, ML model not built"),
          ("Layer 2: Quick Photo",
           "Camera opens with 1 tap. AI identifies food and estimates macros from the photo.",
           "Not started"),
          ("Layer 3: Learned Routine",
           "After 2-3 weeks, recognizes meals by time + location. “Your usual lunch?” — 1 tap.",
           "Not started"),
          ("Layer 4: Biometric Inference",
           "With enough history, AI infers the most likely meal from context alone.",
           "Research phase")]
positions = [(Inches(0.9), Inches(2.35)), (Inches(6.85), Inches(2.35)),
             (Inches(0.9), Inches(4.75)), (Inches(6.85), Inches(4.75))]
cw, chh = Inches(5.55), Inches(2.2)
for (head, body, status), (x, y) in zip(layers, positions):
    card(s, x, y, cw, chh, CARD, CARD)
    tf = textbox(s, x+Inches(0.3), y+Inches(0.2), cw-Inches(0.6), chh-Inches(0.4), MSO_ANCHOR.TOP)
    para(tf, head, 15, bold=True, color=ACCENT, first=True, space_after=6)
    para(tf, body, 12, color=RGBColor(0xAA,0xAA,0xAA), space_after=6, line=1.15)
    para(tf, status, 11, bold=True, color=WARN)

# =====================================================================
# SLIDE 10 — Challenge 2 (4 cards)
# =====================================================================
challenge_slide("Differentiator 2 — The Challenges",
    "The hardest technical challenge in the product",
    [("The ML detection model doesn't exist yet.",
      "Published research achieves 85% precision / 81% recall in controlled conditions. Real-world will be lower. Training requires real eating data from diverse users."),
     ("Frequent false positives.",
      "Brushing teeth, drinking water, scratching your face — all involve hand-to-mouth movement. If the watch asks “Eating now?” while brushing teeth, user trust erodes."),
     ("Watch battery is a real constraint.",
      "Continuous accelerometer + gyroscope monitoring on a coin-sized battery. If VitalAI visibly drains battery, users uninstall."),
     ("Intentional confirmation.",
      "We always ask for a 1-tap confirmation. Never “zero friction” — but keeps users conscious of what they eat and ensures data quality.")])

# =====================================================================
# SLIDE 11 — Differentiator 3: Feedback Loop
# =====================================================================
s = slide()
diff_header(s, "Differentiator 3", "Nutrition Feedback Loop")
tf = textbox(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(0.8), MSO_ANCHOR.MIDDLE)
para(tf, [("Profile → ", TEXT, False), ("Context detected", ACCENT, True),
          (" → Personalized suggestion → ", TEXT, False), ("Feedback", ACCENT, True),
          (" → Learning → ", TEXT, False), ("Better suggestion", GREEN, True)],
     19, PP_ALIGN.CENTER, first=True, line=1.4)
tf2 = textbox(s, Inches(1.5), Inches(4.2), Inches(10.3), Inches(1.2), MSO_ANCHOR.TOP)
para(tf2, "Every rating makes the next suggestion more aligned with the user's real taste.",
     16, PP_ALIGN.CENTER, color=RGBColor(0xAA,0xAA,0xAA), first=True, space_after=10)
para(tf2, [("This creates a ", RGBColor(0xAA,0xAA,0xAA), False),
           ("switching cost", ACCENT, True),
           (" — leaving means starting over from scratch.", RGBColor(0xAA,0xAA,0xAA), False)],
     16, PP_ALIGN.CENTER)

# =====================================================================
# SLIDE 12 — Challenge 3
# =====================================================================
challenge_slide("Differentiator 3 — The Challenges",
    "From “simple” to “intelligent” is a big leap",
    [("Current learning is primitive.",
      "Today, the system only excludes recipes rated ≤ 2 stars. It doesn't understand ingredients, flavors, or contextual patterns. Real personalization requires a recommendation model with ML."),
     ("Cold start problem.",
      "New users have zero history. The first 2-3 weeks of suggestions are essentially random. If those early suggestions are bad, the user churns before the system gets a chance to learn."),
     ("Feedback fatigue.",
      "Asking for a rating after every meal adds friction. Users rate the first few meals enthusiastically, then stop. We need to balance data collection with user experience.")])

# =====================================================================
# SLIDE 13 — Competitive Landscape
# =====================================================================
s = slide()
heading(s, "Competitive Landscape", size=32, top=Inches(0.5))
comp = [("Competitor", "What they do well", "What they lack"),
        ("Oura Advisor", "LLM + ring data", "Reactive, no nutrition, $300 ring"),
        ("Fitbit + Gemini", "Google's scale", "In preview, not proactive, no feedback"),
        ("Apple Health+", "Native ecosystem", "No nutrition, no feedback loop"),
        ("WHOOP Coach", "Athlete recovery", "Proprietary hardware, no nutrition"),
        ("VitalAI", "Proactivity + meal detection + feedback loop", "")]
t = s.shapes.add_table(len(comp), 3, Inches(0.9), Inches(1.55), Inches(11.5), Inches(2.9)).table
t.columns[0].width = Inches(3.0)
t.columns[1].width = Inches(4.6)
t.columns[2].width = Inches(3.9)
for ri, row in enumerate(comp):
    is_vital = (row[0] == "VitalAI")
    for ci, val in enumerate(row):
        cell = t.cell(ri, ci)
        cell.margin_left = Inches(0.15); cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x12,0x22,0x2A) if ri == 0 else BG
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.size = Pt(13); r.font.name = FONT
        r.font.bold = (ri == 0 or is_vital)
        r.font.color.rgb = ACCENT if (ri == 0 or is_vital) else TEXT
    if is_vital:
        t.cell(ri, 1).merge(t.cell(ri, 2))
card(s, Inches(0.9), Inches(4.75), Inches(11.5), Inches(1.6), WARN_BG, WARN)
tf = textbox(s, Inches(1.25), Inches(4.9), Inches(10.9), Inches(1.3), MSO_ANCHOR.MIDDLE)
para(tf, [("Honest risk:  ", WARN, True),
          ("If Apple or Google build proactive coaching natively, they have distribution we can never match. Our defense is the nutrition feedback loop and speed of iteration.", TEXT, False)],
     14, first=True, line=1.2)

# =====================================================================
# SLIDE 14 — Market
# =====================================================================
s = slide()
heading(s, "Market", size=36, top=Inches(0.8))
mstats = [("$12.25B", "Digital health coaching\nglobal (2025)"),
          ("$1.5B", "Digital health\nBrazil (+18% YoY)"),
          ("~23M", "Estimated target\naudience in Brazil")]
for (num, lab), x in zip(mstats, xs):
    tf = textbox(s, x, Inches(2.5), Inches(3.1), Inches(1.0), MSO_ANCHOR.TOP)
    para(tf, num, 44, PP_ALIGN.CENTER, bold=True, color=ACCENT, first=True, space_after=0)
    tf2 = textbox(s, x, Inches(3.7), Inches(3.1), Inches(1.0), MSO_ANCHOR.TOP)
    para(tf2, lab, 14, PP_ALIGN.CENTER, color=MUTED, first=True, line=1.2)
tf = textbox(s, Inches(1.0), Inches(5.4), Inches(11.3), Inches(0.8), MSO_ANCHOR.TOP)
para(tf, "Active young professionals (8M) + Executives focused on longevity (3M) + Women with defined fitness goals (12M)",
     14, PP_ALIGN.CENTER, color=MUTED, first=True)

# =====================================================================
# SLIDE 15 — Business Model
# =====================================================================
s = slide()
heading(s, "Business Model", size=32, top=Inches(0.5))
plans = [("Plan", "What's included", "Price"),
         ("Free", "Morning message + 2 suggestions/day + basic summary", "Free"),
         ("Premium", "Full-day coaching + nutrition feedback + weekly report", "~$8/mo"),
         ("Pro", "Everything + monthly report + PDF for professionals", "~$16/mo")]
t = s.shapes.add_table(len(plans), 3, Inches(0.9), Inches(1.5), Inches(11.5), Inches(2.0)).table
t.columns[0].width = Inches(2.2)
t.columns[1].width = Inches(7.0)
t.columns[2].width = Inches(2.3)
for ri, row in enumerate(plans):
    prem = (row[0] == "Premium")
    for ci, val in enumerate(row):
        cell = t.cell(ri, ci)
        cell.margin_left = Inches(0.15); cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x12,0x22,0x2A) if ri == 0 else BG
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.size = Pt(13); r.font.name = FONT
        r.font.bold = (ri == 0 or (prem and ci == 0))
        r.font.color.rgb = ACCENT if (ri == 0 or (prem and ci == 0)) else TEXT
card(s, Inches(0.9), Inches(3.8), Inches(11.5), Inches(1.4), WARN_BG, WARN)
tf = textbox(s, Inches(1.25), Inches(3.95), Inches(10.9), Inches(1.1), MSO_ANCHOR.MIDDLE)
para(tf, [("Honest unknown:  ", WARN, True),
          ("We have no willingness-to-pay data yet. The Free tier may be “good enough” for most users. The weekly report is our best conversion anchor — but needs validation.", TEXT, False)],
     14, first=True, line=1.2)
tf2 = textbox(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(0.8), MSO_ANCHOR.TOP)
para(tf2, [("B2B potential:  ", ACCENT, True),
           ("Companies offer VitalAI as an employee health benefit — ~$6/employee/month. Could be stronger than B2C.", RGBColor(0xAA,0xAA,0xAA), False)],
     14, first=True, line=1.2)

# =====================================================================
# SLIDE 16 — Regulatory & Privacy
# =====================================================================
challenge_slide("",
    "Regulatory & Privacy",
    [("Positioning: Wellness, not diagnosis.",
      "“Your HRV is low, try breathing” = wellness. “Your heart rate suggests arrhythmia risk” = medical claim (regulatory territory). The line gets blurrier as we add features."),
     ("Data protection (LGPD/GDPR).",
      "Health data is “sensitive data.” Explicit consent, right to deletion, portability required. Architecture designed for this — implementation partial, must complete before launch."),
     ("Privacy by design.",
      "On-device processing — raw biometric data never leaves the device. Only classified events are sent to the backend.")],
    bar=ACCENT, fill=ACC_BG, strong=ACCENT)
# fix the header tag/title for this slide (challenge_slide put subtitle as title)
# (subtitle already serves as the heading "Regulatory & Privacy")

# =====================================================================
# SLIDE 17 — Why Luli + VitalAI
# =====================================================================
s = slide()
tf = textbox(s, Inches(0.9), Inches(0.6), Inches(11.5), Inches(0.8), MSO_ANCHOR.TOP)
para(tf, [("Why ", WHITE, True), ("Luli", ACCENT, True),
          (" + ", WHITE, True), ("VitalAI", ACCENT, True)], 34, PP_ALIGN.CENTER, first=True)
cols = [("What Luli brings",
         ["Real user base to validate hypotheses",
          "Brand trust in the fitness market",
          "Domain expertise — trainers and nutritionists",
          "Existing distribution channel",
          "B2B network for corporate partnerships"]),
        ("What I bring",
         ["A working technical foundation — not slides",
          "Deep understanding of the hard problems",
          "Architecture designed for scale and privacy",
          "Clear-eyed roadmap of what's hard and why",
          "Proven ability to execute"])]
xcol = [Inches(1.1), Inches(7.0)]
for (title, bullets), x in zip(cols, xcol):
    card(s, x, Inches(1.9), Inches(5.2), Inches(4.4), CARD, ACCENT)
    tf = textbox(s, x+Inches(0.4), Inches(2.15), Inches(4.5), Inches(4.0), MSO_ANCHOR.TOP)
    para(tf, title, 18, bold=True, color=ACCENT, first=True, space_after=12)
    for b in bullets:
        para(tf, [("•  ", ACCENT, False), (b, TEXT, False)], 14, space_after=10, line=1.15)

# =====================================================================
# SLIDE 18 — What's Already Built
# =====================================================================
s = slide()
heading(s, "What's Already Built", size=32, top=Inches(0.6))
tf = textbox(s, Inches(0.9), Inches(1.3), Inches(11.5), Inches(0.5), MSO_ANCHOR.TOP)
para(tf, "Not just an idea — a working foundation", 18, PP_ALIGN.CENTER, color=MUTED, first=True)
built = [("100%", "MVP complete"), ("85%", "Alpha (nutrition)"),
         ("11", "API endpoints"), ("6", "health triggers working")]
bx = [Inches(0.9), Inches(3.9), Inches(6.9), Inches(9.9)]
for (num, lab), x in zip(built, bx):
    tf = textbox(s, x, Inches(2.6), Inches(2.6), Inches(0.9), MSO_ANCHOR.TOP)
    para(tf, num, 40, PP_ALIGN.CENTER, bold=True, color=ACCENT, first=True, space_after=0)
    tf2 = textbox(s, x, Inches(3.6), Inches(2.6), Inches(0.8), MSO_ANCHOR.TOP)
    para(tf2, lab, 14, PP_ALIGN.CENTER, color=MUTED, first=True, line=1.15)
tf = textbox(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(0.6), MSO_ANCHOR.TOP)
para(tf, "iOS + Apple Watch + Backend (Node.js/TypeScript) + Claude AI + HealthKit + Push Notifications + Nutrition + Reports",
     13, PP_ALIGN.CENTER, color=DIM, first=True)
tf2 = textbox(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.5), MSO_ANCHOR.TOP)
para(tf2, "Full technical documentation available on request.", 12, PP_ALIGN.CENTER, color=DIM, first=True)

# =====================================================================
# SLIDE 19 — Close
# =====================================================================
s = slide()
tf = textbox(s, Inches(1.2), Inches(2.2), Inches(10.9), Inches(3.0), MSO_ANCHOR.MIDDLE)
para(tf, "I'm not selling a finished product.", 26, PP_ALIGN.CENTER, bold=True, color=WHITE,
     first=True, space_after=18, line=1.3)
para(tf, "I'm showing a technical foundation, a clear-eyed view of the hard problems, and the conviction that together we can solve them.",
     22, PP_ALIGN.CENTER, color=RGBColor(0xDD,0xDD,0xDD), line=1.4)

prs.save("/home/user/VitalAI/docs/luli-fitness-presentation.pptx")
print(f"Saved {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
